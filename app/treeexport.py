"""Export the archive tree — every folder with its subtree's file count and
size, plus totals — as CSV, Excel or PDF.

Reads the local index only (no AWS calls), so it's cheap and works while Deep
Archive objects are cold. Every export carries the bucket name and the local
generation time in its header, so a printed or filed copy is self-describing.

Excel needs `openpyxl` and PDF needs `reportlab`; both are in requirements.txt.
If one is missing the endpoint says so instead of failing obscurely.
"""

import csv
import io
import re
from datetime import datetime

import db
import version


def fmt_bytes(n):
    """Decimal (SI) units, as DIT/offload tools and drive labels report them."""
    n = float(n or 0)
    for u in ("B", "KB", "MB", "GB", "TB", "PB"):
        if n < 1000 or u == "PB":
            return f"{n:.0f} {u}" if u == "B" else f"{n:.1f} {u}"
        n /= 1000


def rollup(bucket, prefix=""):
    """One pass over the index, accumulating each file's size into every one of
    its ancestor folders — so a folder's numbers cover its whole subtree, the
    same totals the Tree tab shows.

    Returns (folders, files, totals).
    """
    files = db.files_for_export(bucket, prefix)
    folders = {}
    total_bytes = 0
    for f in files:
        size = f["size"] or 0
        total_bytes += size
        path = ""
        for part in f["key"].split("/")[:-1]:
            path += part + "/"
            e = folders.setdefault(path, [0, 0])
            e[0] += 1
            e[1] += size
    rows = [{"path": p,
             "name": p.rstrip("/").split("/")[-1],
             "depth": p.rstrip("/").count("/"),
             "count": c,
             "bytes": b} for p, (c, b) in sorted(folders.items())]
    return rows, files, {"folders": len(rows), "files": len(files),
                         "bytes": total_bytes}


def _meta(bucket, prefix, totals):
    return [
        ("Bucket", f"s3://{bucket}"),
        ("Scope", prefix or "(whole bucket)"),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("ColdVault", f"v{version.VERSION}"),
        ("Folders", f"{totals['folders']:,}"),
        ("Files", f"{totals['files']:,}"),
        ("Total size", f"{fmt_bytes(totals['bytes'])} ({totals['bytes']:,} bytes)"),
    ]


# ---- CSV -------------------------------------------------------------------

def to_csv(bucket, prefix):
    folders, files, totals = rollup(bucket, prefix)
    buf = io.StringIO()
    w = csv.writer(buf)
    for k, v in _meta(bucket, prefix, totals):
        w.writerow([k, v])
    w.writerow([])
    w.writerow(["Folder", "Depth", "Files", "Size (bytes)", "Size"])
    for f in folders:
        w.writerow([f["path"], f["depth"], f["count"], f["bytes"],
                    fmt_bytes(f["bytes"])])
    w.writerow([])
    w.writerow(["TOTAL", "", totals["files"], totals["bytes"],
                fmt_bytes(totals["bytes"])])
    w.writerow([])
    w.writerow(["Key", "Size (bytes)", "Size", "Status", "Storage class",
                "Uploaded", "Manifest", "Expires"])
    for f in files:
        w.writerow([f["key"], f["size"] or 0, fmt_bytes(f["size"]), f["status"],
                    f["storage_class"] or "", f["uploaded_at"] or "",
                    f["manifest_state"] or "", f["expires_at"] or ""])
    return buf.getvalue().encode("utf-8-sig")   # BOM so Excel reads UTF-8


# ---- Excel -----------------------------------------------------------------

def to_xlsx(bucket, prefix):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise RuntimeError("Excel export needs the 'openpyxl' package — rebuild "
                           "the image: docker compose up -d --build")

    folders, files, totals = rollup(bucket, prefix)
    wb = Workbook()
    bold = Font(bold=True)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")

    def header_row(ws, labels, row):
        for i, label in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=label)
            c.font = head
            c.fill = fill
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    # --- Tree sheet: folders with subtree totals
    ws = wb.active
    ws.title = "Tree"
    r = 1
    for k, v in _meta(bucket, prefix, totals):
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    header_row(ws, ["Folder", "Depth", "Files", "Size (bytes)", "Size"], r)
    r += 1
    for f in folders:
        # Indent by depth so the sheet reads as a tree, but keep the full path.
        ws.cell(row=r, column=1, value=f["path"]).alignment = \
            Alignment(indent=min(f["depth"], 14))
        ws.cell(row=r, column=2, value=f["depth"])
        ws.cell(row=r, column=3, value=f["count"])
        ws.cell(row=r, column=4, value=f["bytes"]).number_format = "#,##0"
        ws.cell(row=r, column=5, value=fmt_bytes(f["bytes"]))
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = bold
    ws.cell(row=r, column=3, value=totals["files"]).font = bold
    c = ws.cell(row=r, column=4, value=totals["bytes"])
    c.font = bold
    c.number_format = "#,##0"
    ws.cell(row=r, column=5, value=fmt_bytes(totals["bytes"])).font = bold
    for col, width in zip("ABCDE", (80, 8, 10, 16, 12)):
        ws.column_dimensions[col].width = width

    # --- Files sheet: every object, so the workbook is auditable
    fs = wb.create_sheet("Files")
    header_row(fs, ["Key", "Size (bytes)", "Size", "Status", "Storage class",
                    "Uploaded", "Verified", "SHA-256", "Manifest", "Expires"], 1)
    for i, f in enumerate(files, start=2):
        fs.cell(row=i, column=1, value=f["key"])
        fs.cell(row=i, column=2, value=f["size"] or 0).number_format = "#,##0"
        fs.cell(row=i, column=3, value=fmt_bytes(f["size"]))
        fs.cell(row=i, column=4, value=f["status"])
        fs.cell(row=i, column=5, value=f["storage_class"] or "")
        fs.cell(row=i, column=6, value=f["uploaded_at"] or "")
        fs.cell(row=i, column=7, value=f["verified_at"] or "")
        fs.cell(row=i, column=8, value=f["sha256"] or "")
        fs.cell(row=i, column=9, value=f["manifest_state"] or "")
        fs.cell(row=i, column=10, value=f["expires_at"] or "")
    for col, width in zip("ABCDEFGHIJ", (80, 16, 12, 12, 14, 20, 20, 20, 14, 12)):
        fs.column_dimensions[col].width = width
    fs.auto_filter.ref = f"A1:J{max(len(files) + 1, 2)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- PDF -------------------------------------------------------------------

def to_pdf(bucket, prefix):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                        Table, TableStyle)
    except ImportError:
        raise RuntimeError("PDF export needs the 'reportlab' package — rebuild "
                           "the image: docker compose up -d --build")

    folders, files, totals = rollup(bucket, prefix)
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=12 * mm, bottomMargin=14 * mm,
        title=f"ColdVault — {bucket}", author=f"ColdVault v{version.VERSION}")

    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15,
                        leading=19, spaceAfter=2)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5, leading=12,
                         textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10)

    story = [Paragraph("❄ ColdVault archive tree", h1)]
    story.append(Paragraph(" &nbsp;·&nbsp; ".join(
        f"<b>{k}:</b> {_x(v)}" for k, v in _meta(bucket, prefix, totals)), sub))
    story.append(Spacer(1, 6 * mm))

    data = [["Folder", "Files", "Size", "Bytes"]]
    for f in folders:
        indent = "&nbsp;" * (4 * min(f["depth"], 10))
        data.append([Paragraph(f"{indent}{_x(f['name'])}/", cell),
                     f"{f['count']:,}", fmt_bytes(f["bytes"]),
                     f"{f['bytes']:,}"])
    data.append(["TOTAL", f"{totals['files']:,}",
                 fmt_bytes(totals["bytes"]), f"{totals['bytes']:,}"])

    avail = doc.width
    table = Table(data, colWidths=[avail - 110 * mm, 25 * mm, 30 * mm, 55 * mm],
                  repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("LINEBELOW", (0, 0), (-1, -2), 0.25, colors.HexColor("#DDDDDD")),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#1F3A5F")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2),
         [colors.white, colors.HexColor("#F6F8FA")]),
    ]))
    story.append(table)
    if not folders:
        story.append(Paragraph("The index holds no objects for this scope.", sub))

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(14 * mm, 8 * mm, f"s3://{bucket} · generated {stamp}")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm,
                               f"page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _x(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


# ---- dispatch --------------------------------------------------------------

FORMATS = {
    "csv":  (to_csv,  "csv",  "text/csv"),
    "xlsx": (to_xlsx, "xlsx",
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf":  (to_pdf,  "pdf",  "application/pdf"),
}


def build(bucket, prefix="", fmt="xlsx"):
    """-> (bytes, filename, mimetype). Raises RuntimeError for a missing
    optional dependency, ValueError for an unknown format."""
    if fmt not in FORMATS:
        raise ValueError(f"format must be one of: {', '.join(FORMATS)}")
    fn, ext, mime = FORMATS[fmt]
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", f"{bucket}-{prefix}".strip("-/")) or "archive"
    name = f"coldvault-{slug[:80]}-{datetime.now():%Y%m%d-%H%M%S}.{ext}"
    return fn(bucket, prefix), name, mime
